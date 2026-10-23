#!/usr/bin/env python3
"""Master Builder Script for InstaHealth 60-Page Academic Project Report
Generates both InstaHealth_60_Page_Project_Report.docx and InstaHealth_60_Page_Project_Report.md
Strictly adhering to The Apollo University, School of Technology, CSE Department institutional template.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from report_sections.preliminary import (
    TITLE_DATA, DECLARATION_TEXT, BONAFIDE_CERTIFICATE_TEXT,
    ACKNOWLEDGEMENT_PARAGRAPHS, ABSTRACT_TEXT, ABSTRACT_KEYWORDS,
    TOC_ENTRIES, FIGURES_LIST, TABLES_LIST, ABBREVIATIONS_LIST
)
from report_sections.chapter1 import CHAPTER_1_TITLE, SECTIONS as CH1_SECTIONS
from report_sections.chapter2 import CHAPTER_2_TITLE, SECTIONS as CH2_SECTIONS
from report_sections.chapter3 import CHAPTER_3_TITLE, SECTIONS as CH3_SECTIONS
from report_sections.chapter4 import CHAPTER_4_TITLE, SECTIONS as CH4_SECTIONS
from report_sections.chapter5 import CHAPTER_5_TITLE, SECTIONS as CH5_SECTIONS
from report_sections.chapter6 import CHAPTER_6_TITLE, SECTIONS as CH6_SECTIONS
from report_sections.chapter7 import CHAPTER_7_TITLE, SECTIONS as CH7_SECTIONS
from report_sections.references_pubs import (
    REFERENCES_TITLE, REFERENCES_LIST,
    PUBLICATIONS_TITLE, PUBLICATION_SUMMARY
)
from report_sections.tables_data import TABLES

OUTPUT_DOCX = "/Users/jathinreddy/Desktop/InstaHealth/InstaHealth_60_Page_Project_Report.docx"
OUTPUT_MD = "/Users/jathinreddy/Desktop/InstaHealth/InstaHealth_60_Page_Project_Report.md"
WORKSPACE_DIR = "/Users/jathinreddy/Desktop/InstaHealth"

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="94A3B8", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def add_page_number_to_run(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_para(doc, text="", size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0, line_spacing=1.5, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = color
    return p

def add_heading_1(doc, text):
    p = add_para(doc, text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=14, line_spacing=1.5, color=RGBColor(30, 58, 138))
    return p

def add_heading_2(doc, text):
    p = add_para(doc, text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=6, line_spacing=1.5, color=RGBColor(15, 23, 42))
    return p

def add_heading_3(doc, text):
    p = add_para(doc, text, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4, line_spacing=1.5, color=RGBColor(30, 41, 59))
    return p

def add_image_figure(doc, img_path, caption_text, width_in=5.4):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(12)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.bold = True
        r_cap.font.color.rgb = RGBColor(51, 65, 85)
        return p_cap
    else:
        print(f"Warning: Image file not found: {img_path}")
        return None

def add_code_snippet(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Border
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_text.strip())
    r.font.name = "Courier New"
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    # Spacing after code block
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_table_data(doc, table_key):
    tdata = TABLES.get(table_key)
    if not tdata:
        return
    
    # Caption
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(12)
    p_cap.paragraph_format.space_after = Pt(6)
    r_cap = p_cap.add_run(tdata["title"])
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10.5)
    r_cap.bold = True
    r_cap.font.color.rgb = RGBColor(30, 58, 138)
    
    cols = len(tdata["headers"])
    rows = len(tdata["rows"]) + 1
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="94A3B8", sz="4", val="single")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for c_idx, title in enumerate(tdata["headers"]):
        cell = hdr_cells[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(title)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, row in enumerate(tdata["rows"]):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = row_cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx > 0 and len(val) < 25) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.0)
            r.font.color.rgb = RGBColor(15, 23, 42)
            
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(8)

def render_section_content(doc, sec):
    add_heading_2(doc, sec["title"])
    for p_text in sec.get("paragraphs", []):
        add_para(doc, p_text, size=12, space_after=6, line_spacing=1.5)
    
    for sub in sec.get("subsections", []):
        add_heading_3(doc, sub["title"])
        for p_text in sub.get("paragraphs", []):
            if p_text.startswith("Algorithm ") or p_text.startswith("// =="):
                add_code_snippet(doc, p_text)
            elif p_text == "[FIGURE 3.1: GANTT CHART]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_gantt_chart.png"),
                                 "Figure 3.1: InstaHealth 12-Week Comprehensive Project Engineering Gantt Schedule", width_in=5.8)
            elif p_text == "[FIGURE 4.1: SYSTEM ARCHITECTURE]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_3_1_architecture.png"),
                                 "Figure 4.1: InstaHealth Multi-Layered Zero-Cloud On-Device Smartwatch Architecture", width_in=5.8)
            elif p_text == "[FIGURE 4.2: DFD LEVEL 0]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_dfd_level0.png"),
                                 "Figure 4.2: Data Flow Diagram (DFD) Level 0: Context Level System Representation", width_in=5.8)
            elif p_text == "[FIGURE 4.3: DFD LEVEL 1]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_dfd_level1.png"),
                                 "Figure 4.3: Data Flow Diagram (DFD) Level 1: Subsystem Level Functional Decomposition", width_in=5.8)
            elif p_text == "[FIGURE 4.4: DFD LEVEL 2]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_dfd_level2.png"),
                                 "Figure 4.4: Data Flow Diagram (DFD) Level 2: Detailed Process Flow for TinyML Micro-SGD Engine", width_in=5.8)
            elif p_text == "[FIGURE 5.1: CALIBRATION WORKFLOW]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_3_2_workflow.png"),
                                 "Figure 5.1: InstaHealth 48-Hour On-Device Calibration and Lifecycle State Machine Workflow", width_in=5.8)
            elif p_text == "[FIGURE 6.1: DAY 1 CALIBRATION SCREENSHOT]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "screenshot_day1_calibration.png"),
                                 "Figure 6.1: InstaHealth Day 1 & Day 2 Calibration UI: Real-Time Telemetry vs Demographic Benchmarks", width_in=3.6)
            elif p_text == "[FIGURE 6.2: DAY 3 PRESCRIPTION SCREENSHOT]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "screenshot_day3_prescriptions.png"),
                                 "Figure 6.2: InstaHealth Day 3+ Calibrated Prescription UI: Multi-Metric Circular Ring Carousel", width_in=3.6)
            elif p_text == "[FIGURE 6.3: ABLATION CHART]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_ablation_chart.png"),
                                 "Figure 6.3: Empirical Ablation Study: Mean Absolute Error (MAE) Across Physiological Models", width_in=5.8)
            elif p_text == "[FIGURE 6.4: BATTERY POWER PROFILE]":
                add_image_figure(doc, os.path.join(WORKSPACE_DIR, "fig_battery_chart.png"),
                                 "Figure 6.4: Hardware Subsystem Current Draw Profile and 48-Hour Smartwatch Battery Budget", width_in=5.8)
            else:
                add_para(doc, p_text, size=12, space_after=6, line_spacing=1.5)

print("Starting compilation of InstaHealth 60-Page Academic Project Report...")
doc = docx.Document()

# Configure Global Section Margins
section = doc.sections[0]
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.25)
section.right_margin = Inches(1.0)

# Configure Header and Footer
footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_r = footer_p.add_run()
footer_r.font.name = "Times New Roman"
footer_r.font.size = Pt(10)
add_page_number_to_run(footer_r)

# ==============================================================================
# 1. TITLE PAGE
# ==============================================================================
print("Generating Title Page...")
add_para(doc, TITLE_DATA["title"], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18, line_spacing=1.5, color=RGBColor(30, 58, 138))
add_para(doc, TITLE_DATA["subtitle"], size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=6)
add_para(doc, TITLE_DATA["submission_text"], size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=12)

# Student Names
for name, reg in TITLE_DATA["students"]:
    add_para(doc, f"{name} ({reg})", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)

add_para(doc, TITLE_DATA["degree_text"], size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=14, line_spacing=1.3)
add_para(doc, TITLE_DATA["guide_header"], size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
add_para(doc, TITLE_DATA["guide_name"], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)
add_para(doc, TITLE_DATA["guide_title"], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=14)

# Logo
logo_path = os.path.join(WORKSPACE_DIR, "apollo_logo.png")
if os.path.exists(logo_path):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(8)
    p_logo.paragraph_format.space_after = Pt(12)
    p_logo.add_run().add_picture(logo_path, width=Inches(1.8))

add_para(doc, TITLE_DATA["institution"], size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=4, line_spacing=1.3)
add_para(doc, TITLE_DATA["date"], size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=0)
doc.add_page_break()

# ==============================================================================
# 2. DECLARATION
# ==============================================================================
print("Generating Declaration...")
add_heading_1(doc, "DECLARATION")
add_para(doc, DECLARATION_TEXT, size=12, space_before=12, space_after=20, line_spacing=1.5)

add_para(doc, "Signature of the Candidate(s) with names:", size=12, bold=True, space_before=24, space_after=12)
for idx, (name, reg) in enumerate(TITLE_DATA["students"], 1):
    add_para(doc, f"{idx}. {name} ({reg})  ____________________________________", size=11, space_before=6, space_after=6)

add_para(doc, "Place : Chittoor", size=12, space_before=24, space_after=2)
add_para(doc, "Date  : April, 2026", size=12, space_before=2, space_after=0)
doc.add_page_break()

# ==============================================================================
# 3. BONAFIDE CERTIFICATE
# ==============================================================================
print("Generating Bonafide Certificate...")
add_para(doc, "THE APOLLO UNIVERSITY, CHITTOOR\nSCHOOL OF TECHNOLOGY", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=14, line_spacing=1.3)
add_heading_1(doc, "BONAFIDE CERTIFICATE")
add_para(doc, BONAFIDE_CERTIFICATE_TEXT, size=12, space_before=14, space_after=24, line_spacing=1.5)

add_para(doc, "Signature of the Guide\nDr. S. Ramesh Kumar, M.Tech., Ph.D.\nAssociate Professor, Dept. of CSE", size=11, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=24, space_after=36, line_spacing=1.2)

# Program Coordinator and Dean Block
tbl_cert = doc.add_table(rows=1, cols=2)
tbl_cert.alignment = WD_TABLE_ALIGNMENT.CENTER
c1, c2 = tbl_cert.rows[0].cells
p1 = c1.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
p1.paragraph_format.line_spacing = 1.2
r1 = p1.add_run("Program Coordinator\nDept. of Computer Science & Engineering\nSchool of Technology\nThe Apollo University")
r1.font.name = "Times New Roman"
r1.font.size = Pt(10.5)
r1.bold = True

p2 = c2.paragraphs[0]
p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p2.paragraph_format.line_spacing = 1.2
r2 = p2.add_run("Prof. D. Jagadeesan, Ph.D.\nDean, School of Technology\nThe Apollo University")
r2.font.name = "Times New Roman"
r2.font.size = Pt(10.5)
r2.bold = True

add_para(doc, "Internal Examiner: _______________________      External Examiner: _______________________", size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=48, space_after=0)
doc.add_page_break()

# ==============================================================================
# 4. ACKNOWLEDGEMENT
# ==============================================================================
print("Generating Acknowledgement...")
add_heading_1(doc, "ACKNOWLEDGEMENT")
for p_text in ACKNOWLEDGEMENT_PARAGRAPHS:
    add_para(doc, p_text, size=12, space_before=4, space_after=8, line_spacing=1.5)

# Student Names Sign-off
add_para(doc, "Candidates:", size=11, bold=True, space_before=16, space_after=4)
for name, reg in TITLE_DATA["students"]:
    add_para(doc, f"{name} ({reg})", size=11, italic=True, space_before=2, space_after=2)
doc.add_page_break()

# ==============================================================================
# 5. TABLE OF CONTENTS
# ==============================================================================
print("Generating Table of Contents...")
add_heading_1(doc, "TABLE OF CONTENTS")

tbl_toc = doc.add_table(rows=len(TOC_ENTRIES) + 1, cols=3)
tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_toc, color="CBD5E1", sz="2", val="single")

hdr = tbl_toc.rows[0].cells
hdr[0].paragraphs[0].add_run("Chapter / Section Title").bold = True
hdr[1].paragraphs[0].add_run("Section").bold = True
hdr[2].paragraphs[0].add_run("Page No.").bold = True
set_cell_background(hdr[0], "1E3A8A")
set_cell_background(hdr[1], "1E3A8A")
set_cell_background(hdr[2], "1E3A8A")
for h in hdr:
    h.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h.paragraphs[0].runs[0].font.name = "Times New Roman"
    h.paragraphs[0].runs[0].font.size = Pt(10)

for idx, (title, sec_no, p_num) in enumerate(TOC_ENTRIES):
    row_cells = tbl_toc.rows[idx + 1].cells
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for c in row_cells:
        set_cell_background(c, bg)
        set_cell_margins(c, top=40, bottom=40, left=80, right=80)
    
    p0 = row_cells[0].paragraphs[0]
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(title)
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(9.5)
    if "CHAPTER" in title or title in ["DECLARATION", "BONAFIDE CERTIFICATE", "ACKNOWLEDGEMENT", "ABSTRACT", "LIST OF FIGURES", "LIST OF TABLES", "LIST OF ABBREVIATIONS", "REFERENCES", "PUBLICATIONS / CONFERENCES"]:
        r0.bold = True
        
    p1 = row_cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(sec_no)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(9.5)
    
    p2 = row_cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(p_num)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.5)
    if "CHAPTER" in title:
        r2.bold = True

doc.add_page_break()

# ==============================================================================
# 6. ABSTRACT
# ==============================================================================
print("Generating Abstract...")
add_heading_1(doc, "ABSTRACT")
add_para(doc, ABSTRACT_TEXT, size=12, space_before=14, space_after=14, line_spacing=1.5)
add_para(doc, "Keywords: " + ", ".join(ABSTRACT_KEYWORDS), size=11, bold=True, italic=True, space_before=8, space_after=12)
add_para(doc, "Page I", size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=0)
doc.add_page_break()

# ==============================================================================
# 7. LIST OF FIGURES
# ==============================================================================
print("Generating List of Figures...")
add_heading_1(doc, "LIST OF FIGURES")

tbl_lof = doc.add_table(rows=len(FIGURES_LIST) + 1, cols=3)
tbl_lof.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_lof, color="CBD5E1", sz="2", val="single")

hdr = tbl_lof.rows[0].cells
hdr[0].paragraphs[0].add_run("Figure No.").bold = True
hdr[1].paragraphs[0].add_run("Figure Caption / Description").bold = True
hdr[2].paragraphs[0].add_run("Page No.").bold = True
set_cell_background(hdr[0], "1E3A8A")
set_cell_background(hdr[1], "1E3A8A")
set_cell_background(hdr[2], "1E3A8A")
for h in hdr:
    h.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h.paragraphs[0].runs[0].font.name = "Times New Roman"
    h.paragraphs[0].runs[0].font.size = Pt(10)

for idx, (f_no, f_title, p_num) in enumerate(FIGURES_LIST):
    row_cells = tbl_lof.rows[idx + 1].cells
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for c in row_cells:
        set_cell_background(c, bg)
        set_cell_margins(c, top=50, bottom=50, left=80, right=80)
        
    p0 = row_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(f_no)
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(9.5)
    r0.bold = True
    
    p1 = row_cells[1].paragraphs[0]
    r1 = p1.add_run(f_title)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(9.5)
    
    p2 = row_cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(p_num)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.5)

add_para(doc, "Page II", size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=0)
doc.add_page_break()

# ==============================================================================
# 8. LIST OF TABLES
# ==============================================================================
print("Generating List of Tables...")
add_heading_1(doc, "LIST OF TABLES")

tbl_lot = doc.add_table(rows=len(TABLES_LIST) + 1, cols=3)
tbl_lot.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_lot, color="CBD5E1", sz="2", val="single")

hdr = tbl_lot.rows[0].cells
hdr[0].paragraphs[0].add_run("Table No.").bold = True
hdr[1].paragraphs[0].add_run("Table Title / Specification").bold = True
hdr[2].paragraphs[0].add_run("Page No.").bold = True
set_cell_background(hdr[0], "1E3A8A")
set_cell_background(hdr[1], "1E3A8A")
set_cell_background(hdr[2], "1E3A8A")
for h in hdr:
    h.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h.paragraphs[0].runs[0].font.name = "Times New Roman"
    h.paragraphs[0].runs[0].font.size = Pt(10)

for idx, (t_no, t_title, p_num) in enumerate(TABLES_LIST):
    row_cells = tbl_lot.rows[idx + 1].cells
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for c in row_cells:
        set_cell_background(c, bg)
        set_cell_margins(c, top=50, bottom=50, left=80, right=80)
        
    p0 = row_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(t_no)
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(9.5)
    r0.bold = True
    
    p1 = row_cells[1].paragraphs[0]
    r1 = p1.add_run(t_title)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(9.5)
    
    p2 = row_cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(p_num)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.5)

add_para(doc, "Page III", size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=0)
doc.add_page_break()

# ==============================================================================
# 9. LIST OF ABBREVIATIONS
# ==============================================================================
print("Generating List of Abbreviations...")
add_heading_1(doc, "LIST OF ABBREVIATIONS")

tbl_loa = doc.add_table(rows=len(ABBREVIATIONS_LIST) + 1, cols=2)
tbl_loa.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_loa, color="CBD5E1", sz="2", val="single")

hdr = tbl_loa.rows[0].cells
hdr[0].paragraphs[0].add_run("Abbreviation").bold = True
hdr[1].paragraphs[0].add_run("Full Expansion / Terminology").bold = True
set_cell_background(hdr[0], "1E3A8A")
set_cell_background(hdr[1], "1E3A8A")
for h in hdr:
    h.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h.paragraphs[0].runs[0].font.name = "Times New Roman"
    h.paragraphs[0].runs[0].font.size = Pt(10)

for idx, (abbr, full_text) in enumerate(ABBREVIATIONS_LIST):
    row_cells = tbl_loa.rows[idx + 1].cells
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for c in row_cells:
        set_cell_background(c, bg)
        set_cell_margins(c, top=40, bottom=40, left=80, right=80)
        
    p0 = row_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(abbr)
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(9.5)
    r0.bold = True
    
    p1 = row_cells[1].paragraphs[0]
    r1 = p1.add_run(full_text)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(9.5)

add_para(doc, "Page IV", size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=0)
doc.add_page_break()

# ==============================================================================
# CHAPTER 1: INTRODUCTION
# ==============================================================================
print("Generating Chapter 1...")
add_heading_1(doc, CHAPTER_1_TITLE)
for sec in CH1_SECTIONS:
    render_section_content(doc, sec)
doc.add_page_break()

# ==============================================================================
# CHAPTER 2: LITERATURE REVIEW
# ==============================================================================
print("Generating Chapter 2...")
add_heading_1(doc, CHAPTER_2_TITLE)
for sec in CH2_SECTIONS:
    render_section_content(doc, sec)
    if sec["id"] == "2.3":
        add_table_data(doc, "table_2_1")
doc.add_page_break()

# ==============================================================================
# CHAPTER 3: OBJECTIVES AND SCOPE OF WORK
# ==============================================================================
print("Generating Chapter 3...")
add_heading_1(doc, CHAPTER_3_TITLE)
for sec in CH3_SECTIONS:
    render_section_content(doc, sec)
    if sec["id"] == "3.3":
        add_table_data(doc, "table_3_1")
doc.add_page_break()

# ==============================================================================
# CHAPTER 4: SYSTEM ANALYSIS & DESIGN
# ==============================================================================
print("Generating Chapter 4...")
add_heading_1(doc, CHAPTER_4_TITLE)
for sec in CH4_SECTIONS:
    render_section_content(doc, sec)
    if sec["id"] == "4.1":
        add_table_data(doc, "table_4_1")
    elif sec["id"] == "4.3":
        add_table_data(doc, "table_4_2")
    elif sec["id"] == "4.4":
        add_table_data(doc, "table_4_3")
        add_table_data(doc, "table_4_4")
doc.add_page_break()

# ==============================================================================
# CHAPTER 5: IMPLEMENTATION
# ==============================================================================
print("Generating Chapter 5...")
add_heading_1(doc, CHAPTER_5_TITLE)
for sec in CH5_SECTIONS:
    render_section_content(doc, sec)
    if sec["id"] == "5.1":
        add_table_data(doc, "table_5_1")
        add_table_data(doc, "table_5_2")
    elif sec["id"] == "5.2":
        add_table_data(doc, "table_5_3")
doc.add_page_break()

# ==============================================================================
# CHAPTER 6: RESULTS AND DISCUSSION
# ==============================================================================
print("Generating Chapter 6...")
add_heading_1(doc, CHAPTER_6_TITLE)
for sec in CH6_SECTIONS:
    render_section_content(doc, sec)
    if sec["id"] == "6.2":
        add_table_data(doc, "table_6_1")
        add_table_data(doc, "table_6_2")
    elif sec["id"] == "6.3":
        add_table_data(doc, "table_6_3")
        add_table_data(doc, "table_6_4")
        add_table_data(doc, "table_6_5")
doc.add_page_break()

# ==============================================================================
# CHAPTER 7: CONCLUSION AND FUTURE SCOPE
# ==============================================================================
print("Generating Chapter 7...")
add_heading_1(doc, CHAPTER_7_TITLE)
for sec in CH7_SECTIONS:
    render_section_content(doc, sec)
doc.add_page_break()

# ==============================================================================
# REFERENCES
# ==============================================================================
print("Generating References...")
add_heading_1(doc, REFERENCES_TITLE)
add_para(doc, "The references below conform strictly to institutional citation guidelines across all required publication categories (single author, two authors, more than two authors, books, e-books, conference proceedings with and without editors, websites, and patents).", size=11, italic=True, space_before=4, space_after=14)

for ref_text in REFERENCES_LIST:
    add_para(doc, ref_text, size=11, space_before=2, space_after=6, line_spacing=1.25)

doc.add_page_break()

# ==============================================================================
# PUBLICATIONS / CONFERENCES
# ==============================================================================
print("Generating Publications / Conferences...")
add_heading_1(doc, PUBLICATIONS_TITLE)
add_para(doc, "CONFERENCE PUBLICATION PREPRINT DRAFT SUMMARY", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=6)
add_para(doc, "Target Conference: " + PUBLICATION_SUMMARY["conference"], size=11, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=12)

add_para(doc, PUBLICATION_SUMMARY["paper_title"], size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=6, color=RGBColor(30, 58, 138))
add_para(doc, PUBLICATION_SUMMARY["authors"], size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=2)
add_para(doc, PUBLICATION_SUMMARY["affiliations"], size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=14)

add_para(doc, "Abstract—" + PUBLICATION_SUMMARY["abstract"], size=10.5, space_before=4, space_after=8, line_spacing=1.3)
add_para(doc, "Index Terms—" + PUBLICATION_SUMMARY["keywords"], size=10, bold=True, italic=True, space_before=2, space_after=12)

add_heading_2(doc, "I. Introduction & System Motivation")
add_para(doc, PUBLICATION_SUMMARY["section_1_intro_summary"], size=11, space_before=2, space_after=6, line_spacing=1.3)

add_heading_2(doc, "II. Proposed Edge TinyML Architecture & Micro-SGD Optimization")
add_para(doc, PUBLICATION_SUMMARY["section_2_methods_summary"], size=11, space_before=2, space_after=6, line_spacing=1.3)

add_heading_2(doc, "III. Experimental Evaluation & Hardware Benchmark Results")
add_para(doc, PUBLICATION_SUMMARY["section_3_results_summary"], size=11, space_before=2, space_after=6, line_spacing=1.3)

add_heading_2(doc, "IV. Conclusion & Edge Privacy Implications")
add_para(doc, PUBLICATION_SUMMARY["section_4_conclusion_summary"], size=11, space_before=2, space_after=14, line_spacing=1.3)

print(f"Saving compiled Word document to: {OUTPUT_DOCX}")
doc.save(OUTPUT_DOCX)
print("InstaHealth_60_Page_Project_Report.docx generated successfully!")

# ==============================================================================
# GENERATE MATCHING MARKDOWN DOCUMENT
# ==============================================================================
print("Generating matching Markdown document...")

md_lines = []
def md_h1(text): md_lines.append(f"\n# {text}\n")
def md_h2(text): md_lines.append(f"\n## {text}\n")
def md_h3(text): md_lines.append(f"\n### {text}\n")
def md_p(text): md_lines.append(f"\n{text}\n")

md_h1(TITLE_DATA["title"])
md_p(f"**{TITLE_DATA['subtitle']}**  \n*{TITLE_DATA['submission_text']}*")
for name, reg in TITLE_DATA["students"]:
    md_lines.append(f"- **{name}** ({reg})")
md_p(f"\n*{TITLE_DATA['degree_text'].replace(chr(10), ' ')}*\n")
md_p(f"**{TITLE_DATA['guide_header']}**  \n**{TITLE_DATA['guide_name']}**  \n*{TITLE_DATA['guide_title']}*")
md_p(f"**{TITLE_DATA['institution'].replace(chr(10), '  ' + chr(10))}**  \n**{TITLE_DATA['date']}**")
md_lines.append("\n---\n")

md_h1("DECLARATION")
md_p(DECLARATION_TEXT)
md_lines.append("\n**Signature of the Candidate(s) with names:**\n")
for idx, (name, reg) in enumerate(TITLE_DATA["students"], 1):
    md_lines.append(f"{idx}. {name} ({reg})  ____________________________________  \n")
md_p("**Place:** Chittoor  \n**Date:** April, 2026")
md_lines.append("\n---\n")

md_h1("BONAFIDE CERTIFICATE")
md_p(BONAFIDE_CERTIFICATE_TEXT)
md_p("**Dr. S. Ramesh Kumar, M.Tech., Ph.D.**  \nAssociate Professor, Department of Computer Science & Engineering  \n*Signature of the Guide*")
md_p("**Program Coordinator**, Dept. of CSE, School of Technology, The Apollo University  \n**Prof. D. Jagadeesan, Ph.D.**, Dean, School of Technology, The Apollo University")
md_lines.append("\n---\n")

md_h1("ACKNOWLEDGEMENT")
for p in ACKNOWLEDGEMENT_PARAGRAPHS:
    md_p(p)
md_lines.append("\n---\n")

md_h1("ABSTRACT")
md_p(ABSTRACT_TEXT)
md_p(f"**Keywords:** {', '.join(ABSTRACT_KEYWORDS)}")
md_lines.append("\n---\n")

# Tables Helper for Markdown
def md_table(tkey):
    t = TABLES.get(tkey)
    if not t: return
    md_lines.append(f"\n#### {t['title']}\n")
    headers = t["headers"]
    md_lines.append("| " + " | ".join(headers) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
    for row in t["rows"]:
        clean_row = [c.replace("\n", " ") for c in row]
        md_lines.append("| " + " | ".join(clean_row) + " |")
    md_lines.append("\n")

def md_section(sec):
    md_h2(sec["title"])
    for p in sec.get("paragraphs", []):
        md_p(p)
    for sub in sec.get("subsections", []):
        md_h3(sub["title"])
        for p in sub.get("paragraphs", []):
            if p.startswith("Algorithm ") or p.startswith("// =="):
                md_lines.append(f"\n```javascript\n{p}\n```\n")
            elif "[FIGURE " in p:
                tag = p.strip()
                if "3.1" in tag: md_lines.append(f"\n![Figure 3.1 Gantt Chart](fig_gantt_chart.png)\n*Figure 3.1: InstaHealth 12-Week Comprehensive Project Engineering Gantt Schedule*\n")
                elif "4.1" in tag: md_lines.append(f"\n![Figure 4.1 Architecture](fig_3_1_architecture.png)\n*Figure 4.1: InstaHealth Multi-Layered Zero-Cloud On-Device Smartwatch Architecture*\n")
                elif "4.2" in tag: md_lines.append(f"\n![Figure 4.2 DFD Level 0](fig_dfd_level0.png)\n*Figure 4.2: Data Flow Diagram (DFD) Level 0: Context Level System Representation*\n")
                elif "4.3" in tag: md_lines.append(f"\n![Figure 4.3 DFD Level 1](fig_dfd_level1.png)\n*Figure 4.3: Data Flow Diagram (DFD) Level 1: Subsystem Level Functional Decomposition*\n")
                elif "4.4" in tag: md_lines.append(f"\n![Figure 4.4 DFD Level 2](fig_dfd_level2.png)\n*Figure 4.4: Data Flow Diagram (DFD) Level 2: Detailed Process Flow for TinyML Micro-SGD Engine*\n")
                elif "5.1" in tag: md_lines.append(f"\n![Figure 5.1 Workflow](fig_3_2_workflow.png)\n*Figure 5.1: InstaHealth 48-Hour On-Device Calibration and Lifecycle State Machine Workflow*\n")
                elif "6.1" in tag: md_lines.append(f"\n![Figure 6.1 Calibration UI](screenshot_day1_calibration.png)\n*Figure 6.1: InstaHealth Day 1 & Day 2 Calibration UI: Real-Time Telemetry vs Demographic Benchmarks*\n")
                elif "6.2" in tag: md_lines.append(f"\n![Figure 6.2 Prescription UI](screenshot_day3_prescriptions.png)\n*Figure 6.2: InstaHealth Day 3+ Calibrated Prescription UI: Multi-Metric Circular Ring Carousel*\n")
                elif "6.3" in tag: md_lines.append(f"\n![Figure 6.3 Ablation](fig_ablation_chart.png)\n*Figure 6.3: Empirical Ablation Study: Mean Absolute Error (MAE) Across Physiological Models*\n")
                elif "6.4" in tag: md_lines.append(f"\n![Figure 6.4 Battery Profile](fig_battery_chart.png)\n*Figure 6.4: Hardware Subsystem Current Draw Profile and 48-Hour Smartwatch Battery Budget*\n")
            else:
                md_p(p)

md_h1(CHAPTER_1_TITLE)
for s in CH1_SECTIONS: md_section(s)

md_h1(CHAPTER_2_TITLE)
for s in CH2_SECTIONS:
    md_section(s)
    if s["id"] == "2.3": md_table("table_2_1")

md_h1(CHAPTER_3_TITLE)
for s in CH3_SECTIONS:
    md_section(s)
    if s["id"] == "3.3": md_table("table_3_1")

md_h1(CHAPTER_4_TITLE)
for s in CH4_SECTIONS:
    md_section(s)
    if s["id"] == "4.1": md_table("table_4_1")
    elif s["id"] == "4.3": md_table("table_4_2")
    elif s["id"] == "4.4":
        md_table("table_4_3")
        md_table("table_4_4")

md_h1(CHAPTER_5_TITLE)
for s in CH5_SECTIONS:
    md_section(s)
    if s["id"] == "5.1":
        md_table("table_5_1")
        md_table("table_5_2")
    elif s["id"] == "5.2": md_table("table_5_3")

md_h1(CHAPTER_6_TITLE)
for s in CH6_SECTIONS:
    md_section(s)
    if s["id"] == "6.2":
        md_table("table_6_1")
        md_table("table_6_2")
    elif s["id"] == "6.3":
        md_table("table_6_3")
        md_table("table_6_4")
        md_table("table_6_5")

md_h1(CHAPTER_7_TITLE)
for s in CH7_SECTIONS: md_section(s)

md_h1(REFERENCES_TITLE)
for r in REFERENCES_LIST:
    md_lines.append(f"- {r}")

md_h1(PUBLICATIONS_TITLE)
md_p(f"### {PUBLICATION_SUMMARY['paper_title']}")
md_p(f"**Target Venue:** {PUBLICATION_SUMMARY['conference']}  \n**Authors:** {PUBLICATION_SUMMARY['authors']}  \n*{PUBLICATION_SUMMARY['affiliations']}*")
md_p(f"**Abstract:** {PUBLICATION_SUMMARY['abstract']}")
md_p(f"**Keywords:** {PUBLICATION_SUMMARY['keywords']}")
md_p(f"**I. Introduction:** {PUBLICATION_SUMMARY['section_1_intro_summary']}")
md_p(f"**II. Methodology:** {PUBLICATION_SUMMARY['section_2_methods_summary']}")
md_p(f"**III. Results:** {PUBLICATION_SUMMARY['section_3_results_summary']}")
md_p(f"**IV. Conclusion:** {PUBLICATION_SUMMARY['section_4_conclusion_summary']}")

with open(OUTPUT_MD, "w", encoding="utf-8") as f:
    f.write("\n".join(md_lines))

print(f"InstaHealth_60_Page_Project_Report.md generated successfully!")
print("Report assembly completed successfully.")
